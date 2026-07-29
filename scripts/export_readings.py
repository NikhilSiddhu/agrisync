"""Export readings from MongoDB to CSV and XLSX.

    python -m scripts.export_readings                    # everything, CSV
    python -m scripts.export_readings --days 7           # last week only
    python -m scripts.export_readings --xlsx             # add a spreadsheet
    python -m scripts.export_readings --summary-only     # just the quality report

Two shapes come out, because looking at data and modelling data want
different things:

  readings_raw.csv     one row per reading, exactly what the box recorded.
                       The source of truth. Never edit it.

  readings_hourly.csv  one row per hour. This is the modelling table.
                       Disease models count WET HOURS, not wet readings,
                       so the raw file has to be collapsed to hours before
                       any of it means anything.

The quality report is the important part. A model fitted on data with
unnoticed gaps produces confident nonsense, and gaps are invisible until
you go looking for them.
"""
import argparse
import csv
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pymongo import MongoClient

from src.config import Config

WET_THRESHOLD = 50.0          # leaf_wetness at or above this counts as wet
EXPECTED_INTERVAL_S = 60      # firmware POST_INTERVAL_MS

RAW_COLUMNS = [
    "ts", "device_id",
    "air_temp_c", "humidity_pct", "pressure_hpa", "dew_point_c",
    "leaf_temp_c", "leaf_temp_raw_c", "mlx_ambient_c", "leaf_temp_method",
    "leaf_wetness", "leaf_wetness_source",
    "rain_analog", "rain_tips",
    "soil_moisture_pct", "soil_temp_c", "voc_index",
    "battery_v", "rssi",
    "ts_source", "backfilled", "delay_seconds",
    "firmware", "buffered_count",
    "sensors_ok", "quality_flags",
]

HOURLY_COLUMNS = [
    "hour_utc", "hour_local", "local_date", "local_hour", "device_id",
    "n_readings", "coverage_pct",
    "air_temp_mean", "air_temp_min", "air_temp_max",
    "humidity_mean", "dew_point_mean",
    "leaf_temp_mean", "leaf_air_delta_mean",
    "wetness_mean", "wet_fraction", "wet_minutes", "is_wet_hour",
    "rain_tips_in_hour", "rain_analog_min",
    "battery_v_mean", "rssi_mean",
    "n_suspect", "usable",
]


def fetch(db, days=None, device=None, clean_only=False):
    query = {}
    if device:
        query["device_id"] = device
    if days:
        query["ts"] = {"$gte": datetime.now(timezone.utc) - timedelta(days=days)}
    if clean_only:
        query["quality_flags"] = {"$exists": False}
    return list(db.readings.find(query).sort("ts", 1))


def write_raw(rows, path):
    with open(path, "w", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=RAW_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for doc in rows:
            out = {k: doc.get(k) for k in RAW_COLUMNS}
            ts = doc.get("ts")
            out["ts"] = ts.isoformat() if hasattr(ts, "isoformat") else ts
            sensors = doc.get("sensors") or {}
            out["sensors_ok"] = ";".join(k for k, v in sorted(sensors.items()) if v)
            out["quality_flags"] = ";".join(doc.get("quality_flags") or [])
            writer.writerow(out)
    return len(rows)


def _local(dt):
    """Site local time. Everything is stored in UTC, which is correct, but
    UTC hours are unreadable when you are trying to reason about 'last
    night'. Hapur is UTC+5:30, so a wet night straddles two UTC dates.

    The tzinfo line below is not optional. PyMongo hands back NAIVE
    datetimes even though Mongo stores UTC, and astimezone() on a naive
    value assumes system local time. The Pi's clock is set to IST, so
    without this it would read a UTC timestamp as IST and shift nothing -
    every local time silently wrong by five and a half hours.
    """
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    try:
        from zoneinfo import ZoneInfo
        return dt.astimezone(ZoneInfo(Config.SITE_TZ))
    except Exception:
        return dt + timedelta(hours=5, minutes=30)


def _mean(values):
    vals = [v for v in values if v is not None]
    return round(sum(vals) / len(vals), 3) if vals else None


def aggregate_hourly(rows):
    buckets = defaultdict(list)
    for doc in rows:
        ts = doc.get("ts")
        if not hasattr(ts, "replace"):
            continue
        hour = ts.replace(minute=0, second=0, microsecond=0)
        buckets[(hour, doc.get("device_id"))].append(doc)

    expected = 3600 // EXPECTED_INTERVAL_S
    out = []

    for (hour, device), docs in sorted(buckets.items()):
        clean = [d for d in docs if not d.get("quality_flags")]
        source = clean or docs   # fall back rather than emitting nothing

        wetness = [d["leaf_wetness"] for d in source if d.get("leaf_wetness") is not None]
        wet_count = sum(1 for w in wetness if w >= WET_THRESHOLD)
        wet_fraction = (wet_count / len(wetness)) if wetness else None

        air = [d.get("air_temp_c") for d in source]
        leaf = [d.get("leaf_temp_c") for d in source]
        deltas = [
            d["leaf_temp_c"] - d["air_temp_c"]
            for d in source
            if d.get("leaf_temp_c") is not None and d.get("air_temp_c") is not None
        ]

        tips = [d.get("rain_tips") for d in source if d.get("rain_tips") is not None]
        rain_in_hour = (max(tips) - min(tips)) if len(tips) >= 2 else 0

        rain_analog = [d.get("rain_analog") for d in source if d.get("rain_analog") is not None]
        coverage = round(100.0 * len(docs) / expected, 1)

        local = _local(hour)
        out.append({
            "hour_utc": hour.isoformat(),
            "hour_local": local.strftime("%Y-%m-%d %H:%M"),
            "local_date": local.strftime("%Y-%m-%d"),
            "local_hour": local.hour,
            "device_id": device,
            "n_readings": len(docs),
            "coverage_pct": coverage,
            "air_temp_mean": _mean(air),
            "air_temp_min": min((v for v in air if v is not None), default=None),
            "air_temp_max": max((v for v in air if v is not None), default=None),
            "humidity_mean": _mean([d.get("humidity_pct") for d in source]),
            "dew_point_mean": _mean([d.get("dew_point_c") for d in source]),
            "leaf_temp_mean": _mean(leaf),
            "leaf_air_delta_mean": _mean(deltas),
            "wetness_mean": _mean(wetness),
            "wet_fraction": round(wet_fraction, 3) if wet_fraction is not None else None,
            # Fraction of the hour the leaf was wet. This is the number the
            # disease models actually consume.
            "wet_minutes": round(wet_fraction * 60, 1) if wet_fraction is not None else None,
            "is_wet_hour": (1 if wet_fraction is not None and wet_fraction >= 0.5 else 0),
            "rain_tips_in_hour": rain_in_hour,
            "rain_analog_min": min(rain_analog) if rain_analog else None,
            "battery_v_mean": _mean([d.get("battery_v") for d in source]),
            "rssi_mean": _mean([d.get("rssi") for d in source]),
            "n_suspect": len(docs) - len(clean),
            # An hour with under half its readings cannot be trusted for
            # wet-hour accounting. Flagged rather than dropped.
            "usable": 1 if coverage >= 50 else 0,
        })
    return out


def write_hourly(hours, path):
    with open(path, "w", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=HOURLY_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(hours)
    return len(hours)


def quality_report(rows, hours):
    if not rows:
        return ["No readings found."]

    first, last = rows[0]["ts"], rows[-1]["ts"]
    span_s = (last - first).total_seconds()
    expected = int(span_s / EXPECTED_INTERVAL_S) + 1
    actual = len(rows)

    # Longest silence between consecutive readings
    worst_gap, worst_at = timedelta(0), None
    for prev, curr in zip(rows, rows[1:]):
        gap = curr["ts"] - prev["ts"]
        if gap > worst_gap:
            worst_gap, worst_at = gap, prev["ts"]

    suspect = sum(1 for d in rows if d.get("quality_flags"))
    backfilled = sum(1 for d in rows if d.get("backfilled"))
    device_ts = sum(1 for d in rows if d.get("ts_source") == "device")
    with_wetness = sum(1 for d in rows if d.get("leaf_wetness") is not None)
    usable_hours = sum(1 for h in hours if h["usable"])
    wet_hours = sum(1 for h in hours if h["is_wet_hour"])

    flag_counts = defaultdict(int)
    for doc in rows:
        for flag in doc.get("quality_flags") or []:
            flag_counts[flag] += 1

    lines = [
        "DATA QUALITY REPORT",
        "=" * 58,
        f"  Window          {first:%Y-%m-%d %H:%M} to {last:%Y-%m-%d %H:%M} UTC",
        f"  Duration        {span_s / 3600:.1f} hours",
        f"  Readings        {actual} of {expected} expected",
        f"  Coverage        {100.0 * actual / expected:.1f}%",
        f"  Longest gap     {worst_gap.total_seconds() / 60:.0f} minutes"
        + (f" starting {worst_at:%Y-%m-%d %H:%M}" if worst_at else ""),
        "",
        f"  Suspect rows    {suspect} ({100.0 * suspect / actual:.1f}%)",
        f"  Backfilled      {backfilled} (recovered from the board's buffer)",
        f"  Device clock    {device_ts} of {actual} stamped by the board's own NTP",
        f"  With wetness    {with_wetness} ({100.0 * with_wetness / actual:.1f}%)",
        "",
        f"  Hours total     {len(hours)}",
        f"  Hours usable    {usable_hours} (at least 50% coverage)",
        f"  Wet hours       {wet_hours}",
    ]

    if flag_counts:
        lines.append("")
        lines.append("  Quality flags seen:")
        for flag, count in sorted(flag_counts.items(), key=lambda kv: -kv[1]):
            lines.append(f"    {count:>6}  {flag}")

    lines.append("")
    lines.append("  Read this before trusting any model output:")
    if actual / expected < 0.9:
        lines.append("    ! Coverage below 90%. Wet-hour counts will be underestimates.")
    if worst_gap > timedelta(hours=2):
        lines.append("    ! A gap over 2 hours. If it fell on a wet night, that night is lost.")
    if with_wetness < actual * 0.9:
        lines.append("    ! Many rows lack leaf_wetness. MLX90614 or humidity was down.")
    if suspect > actual * 0.05:
        lines.append("    ! Over 5% suspect. Check the sensor before fitting anything.")
    if (actual / expected >= 0.9 and worst_gap <= timedelta(hours=2)
            and with_wetness >= actual * 0.9 and suspect <= actual * 0.05):
        lines.append("    Looks sound. No coverage or quality problems detected.")

    return lines


def write_xlsx(rows, hours, path, report):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed - skipping XLSX. Fix with:")
        print("  pip install openpyxl")
        return False

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="44546A")

    def sheet(ws, columns, records):
        ws.append(columns)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for rec in records:
            ws.append([rec.get(c) for c in columns])
        ws.freeze_panes = "A2"
        for i, col in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(11, min(len(col) + 3, 22))
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Quality report"
    for line in report:
        ws.append([line])
    ws.column_dimensions["A"].width = 70
    ws["A1"].font = Font(bold=True, size=13)

    raw_records = []
    for doc in rows:
        rec = {k: doc.get(k) for k in RAW_COLUMNS}
        ts = doc.get("ts")
        rec["ts"] = ts.replace(tzinfo=None) if hasattr(ts, "replace") else ts
        sensors = doc.get("sensors") or {}
        rec["sensors_ok"] = ";".join(k for k, v in sorted(sensors.items()) if v)
        rec["quality_flags"] = ";".join(doc.get("quality_flags") or [])
        raw_records.append(rec)

    sheet(wb.create_sheet("Hourly"), HOURLY_COLUMNS, hours)
    sheet(wb.create_sheet("Raw readings"), RAW_COLUMNS, raw_records)

    wb.save(path)
    return True


def main():
    ap = argparse.ArgumentParser(description="Export AgriSync readings")
    ap.add_argument("--days", type=float, help="only the last N days")
    ap.add_argument("--device", help="filter to one device_id")
    ap.add_argument("--clean-only", action="store_true",
                    help="drop rows that failed plausibility checks")
    ap.add_argument("--xlsx", action="store_true", help="also write a spreadsheet")
    ap.add_argument("--summary-only", action="store_true", help="print the report, write nothing")
    ap.add_argument("--out", default="exports", help="output directory")
    args = ap.parse_args()

    client = MongoClient(Config.MONGO_URI, serverSelectionTimeoutMS=5000)
    rows = fetch(client.get_default_database(), args.days, args.device, args.clean_only)

    if not rows:
        print("No readings matched. Is the device reporting?")
        return 1

    hours = aggregate_hourly(rows)
    report = quality_report(rows, hours)
    print("\n".join(report))

    if args.summary_only:
        return 0

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M")

    raw_path = out_dir / f"readings_raw_{stamp}.csv"
    hourly_path = out_dir / f"readings_hourly_{stamp}.csv"
    write_raw(rows, raw_path)
    write_hourly(hours, hourly_path)

    print()
    print(f"  {raw_path}     {len(rows)} rows")
    print(f"  {hourly_path}  {len(hours)} rows  <- use this one for modelling")

    if args.xlsx:
        xlsx_path = out_dir / f"agrisync_{stamp}.xlsx"
        if write_xlsx(rows, hours, xlsx_path, report):
            print(f"  {xlsx_path}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
